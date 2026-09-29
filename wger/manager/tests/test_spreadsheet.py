# This file is part of wger Workout Manager.
#
# wger Workout Manager is free software: you can redistribute it and/or modify
# it under the terms of the GNU Affero General Public License as published by
# the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version.
#
# wger Workout Manager is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU General Public License for more details.
#
# You should have received a copy of the GNU Affero General Public License
# along with Workout Manager.  If not, see <http://www.gnu.org/licenses/>.

# Standard Library
import csv
import io
import zipfile
from unittest.mock import patch

# Django
from django.core.files.uploadedfile import SimpleUploadedFile

# Third Party
import openpyxl
from rest_framework import status

# wger
from wger.core.tests.api_base_test import ApiBaseTestCase
from wger.core.tests.base_testcase import BaseTestCase
from wger.manager import spreadsheet
from wger.manager.api.serializers import RoutineStructureSerializer
from wger.manager.models import (
    Day,
    RepetitionsConfig,
    Routine,
    Slot,
    SlotEntry,
    WeightConfig,
    WorkoutLog,
    WorkoutSession,
)
from wger.manager.spreadsheet import COLUMNS


PREVIEW = '/api/v2/routine/import-preview/'
CONFIRM = '/api/v2/routine/import-confirm/'

ROUTINE = {
    'routine_name': '@Synthetic Upper Lower',
    'routine_description': '-sample only, not a real plan',
    'start': '2026-10-05',
    'end': '2026-12-27',
}


def synthetic_rows():
    """Synthetic plan: superset, timed set, lb/body weight, ranges, RiR, rest day"""
    upper = {'day_order': '1', 'day_name': 'Upper A', 'day_is_rest': 'no'}
    lower = {'day_order': '3', 'day_name': 'Lower A', 'day_is_rest': 'no'}
    return [
        {
            **ROUTINE,
            **upper,
            'slot_order': '1',
            'entry_order': '1',
            'exercise_name': 'An exercise',
            'sets': '4',
            'reps': '6',
            'max_reps': '8',
            'weight': '60',
            'weight_unit': 'kg',
            'rir': '2',
            'rest': '120',
            'max_rest': '180',
            'notes': '=HYPERLINK("http://example.com")',
        },
        {
            **upper,
            'slot_order': '2',
            'slot_comment': 'Superset',
            'entry_order': '1',
            'exercise_name': 'Very cool exercise',
            'sets': '3',
            'reps': '8',
            'max_reps': '10',
            'weight_unit': 'Body Weight',
        },
        {
            **upper,
            'slot_order': '2',
            'slot_comment': 'Superset',
            'entry_order': '2',
            'exercise_name': 'Boring exercise',
            'sets': '3',
            'reps': '8',
            'weight': '22.5',
            'weight_unit': 'lb',
            'rest': '60',
        },
        {
            **upper,
            'slot_order': '3',
            'entry_order': '1',
            'exercise_name': 'Pending exercise',
            'set_type': 'iso',
            'sets': '3',
            'reps': '45',
            'rep_unit': 'Seconds',
            'notes': "'quoted note",
        },
        {'day_order': '2', 'day_name': 'Rest', 'day_is_rest': 'yes'},
        {
            **lower,
            'slot_order': '1',
            'entry_order': '1',
            'exercise_name': 'a different name',
            'sets': '5',
            'reps': '5',
            'weight': '135',
            'weight_unit': 'lb',
            'rir': '1',
            'max_rir': '2',
        },
    ]


def csv_file(rows, name='plan.csv'):
    """The rows as a CSV the way a spreadsheet program saves them"""
    grid = [[row.get(column) for column in COLUMNS] for row in rows]
    return SimpleUploadedFile(name, spreadsheet.write_csv(grid))


def history():
    """Every logged workout and set, to prove imports never touch them"""
    return (
        sorted(WorkoutSession.objects.values_list(), key=str),
        sorted(WorkoutLog.objects.values_list(), key=str),
    )


def plan_counts():
    return [m.objects.count() for m in (Routine, Day, Slot, SlotEntry, *spreadsheet.PLAN_MODELS)]


def download_rows(response):
    """Rows of an exported file, CSV or XLSX, as the importer reads them"""
    name = 'x.xlsx' if 'spreadsheetml' in response['Content-Type'] else 'x.csv'
    rows, errors = spreadsheet.read_file(SimpleUploadedFile(name, response.content))
    assert not errors, errors
    return [{k: v for k, v in row.items() if k != '_row'} for row in rows]


class SpreadsheetApiTestCase(BaseTestCase, ApiBaseTestCase):
    """Import and export of routine plans as CSV and XLSX"""

    def setUp(self):
        super().setUp()
        self.authenticate('admin')

    def post(self, url, upload, **data):
        upload.seek(0)
        return self.client.post(url, {'file': upload, **data}, format='multipart')

    def import_plan(self, upload, mode='create', expect=status.HTTP_201_CREATED, **data):
        """Preview and confirm, returns the routine"""
        if mode == 'update':
            data['routine'] = data['routine']
        preview = self.post(PREVIEW, upload, mode=mode, **data)
        self.assertEqual(preview.status_code, status.HTTP_200_OK)
        self.assertTrue(preview.data['ok'], preview.data['errors'])
        confirm = self.post(CONFIRM, upload, mode=mode, plan_hash=preview.data['plan_hash'], **data)
        self.assertEqual(confirm.status_code, expect, confirm.data)
        return Routine.objects.get(pk=confirm.data['id'])

    def export(self, routine, kind='csv'):
        response = self.client.get(f'/api/v2/routine/{routine.pk}/export/?file={kind}')
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        return response

    def preview_errors(self, rows, mode='create', **data):
        response = self.post(PREVIEW, csv_file(rows), mode=mode, **data)
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertFalse(response.data['ok'])
        return {(e['row'], e['column']) for e in response.data['errors']}

    #
    # Round trip
    #
    def test_create_imports_the_whole_plan(self):
        before = history()
        routine = self.import_plan(csv_file(synthetic_rows()))

        self.assertEqual(routine.name, '@Synthetic Upper Lower')
        self.assertEqual(routine.description, '-sample only, not a real plan')
        self.assertFalse(routine.is_template)
        days = list(routine.days.order_by('order'))
        self.assertEqual(
            [(d.order, d.name, d.is_rest) for d in days],
            [(1, 'Upper A', False), (2, 'Rest', True), (3, 'Lower A', False)],
        )
        superset = days[0].slots.get(order=2)
        self.assertEqual(superset.comment, 'Superset')
        self.assertEqual([e.exercise_id for e in superset.entries.all()], [2, 3])
        timed = days[0].slots.get(order=3).entries.get()
        self.assertEqual((timed.type, timed.repetition_unit.name), ('iso', 'Seconds'))
        self.assertEqual(timed.comment, "'quoted note")
        first = days[0].slots.get(order=1).entries.get()
        self.assertEqual(first.comment, '=HYPERLINK("http://example.com")')
        self.assertEqual(first.maxrestconfig_set.get(iteration=1).value, 180)
        alias = days[2].slots.get().entries.get()
        self.assertEqual((alias.exercise_id, alias.weight_unit.name), (1, 'lb'))
        self.assertEqual(alias.maxrirconfig_set.get().value, 2)
        self.assertEqual(history(), before)

    def test_export_update_round_trip_changes_nothing(self):
        routine = self.import_plan(csv_file(synthetic_rows()))
        structure = RoutineStructureSerializer(routine).data

        hashes = []
        for kind in ('csv', 'xlsx'):
            upload = SimpleUploadedFile(f'routine.{kind}', self.export(routine, kind).content)
            preview = self.post(PREVIEW, upload, mode='update', routine=routine.pk).data
            self.assertTrue(preview['ok'], preview['errors'])
            self.assertEqual(preview['diff']['create'], {'days': 0, 'slots': 0, 'entries': 0})
            self.assertEqual(preview['diff']['delete'], {'days': 0, 'slots': 0, 'entries': 0})
            hashes.append(preview['plan_hash'])
        # The hash is of the resolved plan, not the file bytes
        self.assertEqual(hashes[0], hashes[1])

        self.import_plan(upload, 'update', status.HTTP_200_OK, routine=routine.pk)
        routine.refresh_from_db()
        self.assertEqual(RoutineStructureSerializer(routine).data, structure)

    def test_export_create_export_gives_the_same_rows(self):
        routine = self.import_plan(csv_file(synthetic_rows()))
        rows = download_rows(self.export(routine))

        for row in rows:
            row['routine_name'] = 'Copy' if row['routine_name'] else ''
            for column in ('routine_id', 'day_id', 'slot_id', 'entry_id'):
                row[column] = ''
        copy = self.import_plan(csv_file(rows))

        copied = download_rows(self.export(copy))
        for row in copied:
            for column in ('routine_id', 'day_id', 'slot_id', 'entry_id'):
                row[column] = ''
        self.assertEqual(copied, rows)

    #
    # Formula safety
    #
    def test_export_never_writes_formulas(self):
        routine = self.import_plan(csv_file(synthetic_rows()))

        text = self.export(routine).content.decode('utf-8-sig')
        cells = [cell for row in csv.reader(io.StringIO(text)) for cell in row]
        self.assertIn('\'=HYPERLINK("http://example.com")', cells)
        self.assertIn("'@Synthetic Upper Lower", cells)
        self.assertIn("''quoted note", cells)
        self.assertFalse([c for c in cells if c[:1] in ('=', '+', '-', '@')])

        workbook = openpyxl.load_workbook(io.BytesIO(self.export(routine, 'xlsx').content))
        sheet = workbook['routine']
        values = {}
        for row in sheet.iter_rows():
            for cell in row:
                self.assertNotEqual(cell.data_type, 'f')
                if isinstance(cell.value, str):
                    self.assertEqual(cell.data_type, 's')
                    values[cell.value] = cell
        self.assertIn('\'=HYPERLINK("http://example.com")', values)
        self.assertIsInstance(sheet.cell(2, COLUMNS.index('sets') + 1).value, int)

    def test_unescaped_formula_is_refused(self):
        rows = synthetic_rows()
        rows[0].update(notes='=1+1', routine_name='Plain', routine_description='Plain')
        grid = [COLUMNS] + [[row.get(c) or '' for c in COLUMNS] for row in rows]
        out = io.StringIO()
        csv.writer(out).writerows(grid)
        upload = SimpleUploadedFile('plan.csv', out.getvalue().encode())

        errors = self.post(PREVIEW, upload, mode='create').data['errors']
        self.assertEqual([(e['row'], e['column']) for e in errors], [(2, 'notes')])

    #
    # Unsupported structure
    #
    def test_unsupported_structure_is_flagged_and_kept(self):
        routine = self.import_plan(csv_file(synthetic_rows()))
        entry = SlotEntry.objects.filter(slot__day__routine=routine).first()
        progression = WeightConfig.objects.create(
            slot_entry=entry, iteration=2, value=5, operation='+', step='abs', repeat=True
        )
        RepetitionsConfig.objects.filter(slot_entry=entry).update(
            requirements={'rules': ['weight']}
        )
        rows = download_rows(self.export(routine))
        self.assertEqual(rows[0]['unsupported'], 'progression;requirements')

        # Create needs an explicit drop
        for row in rows:
            row['routine_name'] = 'Copy' if row['routine_name'] else ''
        for row in rows:
            for column in ('routine_id', 'day_id', 'slot_id', 'entry_id'):
                row[column] = ''
        self.assertIn((2, 'unsupported'), self.preview_errors(rows))
        preview = self.post(PREVIEW, csv_file(rows), mode='create', drop_unsupported='true').data
        self.assertTrue(preview['ok'], preview['errors'])
        self.assertEqual(
            preview['unsupported_dropped'], [{'row': 2, 'codes': ['progression', 'requirements']}]
        )

        # Update keeps the progression and requirements as they are
        rows = download_rows(self.export(routine))
        rows[0]['weight'] = '65'
        before = list(WeightConfig.objects.filter(iteration__gt=1).values_list())
        self.import_plan(csv_file(rows), 'update', status.HTTP_200_OK, routine=routine.pk)
        self.assertEqual(list(WeightConfig.objects.filter(iteration__gt=1).values_list()), before)
        self.assertEqual(entry.weightconfig_set.get(iteration=1).value, 65)
        self.assertEqual(entry.repetitionsconfig_set.get().requirements, {'rules': ['weight']})
        progression.refresh_from_db()

        # A base value with a gated progression can't be cleared silently
        rows[0]['reps'] = ''
        self.assertIn((2, 'reps'), self.preview_errors(rows, 'update', routine=routine.pk))

    #
    # Exercise identity
    #
    def test_exercise_identity(self):
        cases = [
            ({'exercise_id': '3'}, {'id': 3, 'how': 'id'}),
            ({'exercise_uuid': 'ae3328ba-9a35-4731-bc23-5da50720c5aa'}, {'id': 2, 'how': 'uuid'}),
            ({'exercise_name': 'boring EXERCISE'}, {'id': 3, 'how': 'name'}),
            ({'exercise_name': 'yet another name'}, {'id': 2, 'how': 'name'}),
            ({'exercise_id': '1', 'exercise_name': 'a different name'}, {'id': 1, 'how': 'id'}),
        ]
        for columns, expected in cases:
            rows = synthetic_rows()[:1]
            rows[0].update({'exercise_name': '', **columns})
            data = self.post(PREVIEW, csv_file(rows), mode='create').data
            self.assertTrue(data['ok'], (columns, data['errors']))
            self.assertLessEqual(expected.items(), data['rows'][0]['exercise'].items())

        blocked = [
            ({'exercise_name': 'No such exercise'}, 'unresolved'),
            ({'exercise_name': 'Needed for demo user'}, 'ambiguous'),
            ({'exercise_id': '1', 'exercise_name': 'Boring exercise'}, 'mismatch'),
            ({'exercise_id': '999'}, 'unresolved'),
        ]
        for columns, how in blocked:
            rows = synthetic_rows()[:1]
            rows[0].update({'exercise_name': '', **columns})
            data = self.post(PREVIEW, csv_file(rows), mode='create').data
            self.assertFalse(data['ok'])
            self.assertEqual(data['rows'][0]['exercise']['how'], how)
            self.assertEqual(data['rows'][0]['status'], 'error')
        self.assertEqual([c['id'] for c in data['rows'][0]['exercise'].get('candidates', [])], [])
        rows = synthetic_rows()[:1]
        rows[0]['exercise_name'] = 'Needed for demo user'
        exercise = self.post(PREVIEW, csv_file(rows), mode='create').data['rows'][0]['exercise']
        self.assertEqual([c['id'] for c in exercise['candidates']], [5, 6, 7, 8])

    #
    # Malformed input
    #
    def test_malformed_rows(self):
        def change(index, **values):
            rows = synthetic_rows()
            rows[index].update(values)
            return rows

        superset = synthetic_rows()
        superset[2]['slot_comment'] = 'Other'
        cases = [
            (change(0, reps='abc'), (2, 'reps')),
            (change(0, max_reps='5'), (2, 'max_reps')),
            (change(0, rir='0.3'), (2, 'rir')),
            (change(0, rest='2000'), (2, 'rest')),
            (change(0, sets='0'), (2, 'sets')),
            (change(0, sets=''), (2, 'sets')),
            (change(1, day_name='Other'), (3, 'day_name')),
            (superset, (4, 'slot_comment')),
            (change(2, entry_order='1'), (4, 'entry_order')),
            (change(4, exercise_name='Boring exercise'), (6, 'day_is_rest')),
            (change(0, routine_name='x' * 26), (2, 'routine_name')),
            (change(0, end='2027-03-01'), (2, 'end')),
            (change(0, rep_unit='Parsecs'), (2, 'rep_unit')),
            (change(0, set_type='bouncy'), (2, 'set_type')),
            (change(0, entry_id='1'), (2, 'entry_id')),
            (change(0, day_is_rest='maybe'), (2, 'day_is_rest')),
        ]
        for rows, expected in cases:
            self.assertIn(expected, self.preview_errors(rows))

    def test_file_limits(self):
        def read(name, data):
            with self.assertRaises(spreadsheet.SpreadsheetError) as context:
                spreadsheet.read_file(SimpleUploadedFile(name, data))
            return str(context.exception)

        header = ','.join(COLUMNS)
        self.assertIn(
            'missing columns: unsupported', read('a.csv', header[: -len(',unsupported')].encode())
        )
        self.assertIn('unknown columns: colour', read('a.csv', f'{header},colour'.encode()))
        row = ',' * 6 + '1' + ',' * 25
        self.assertIn('1000 rows', read('a.csv', '\n'.join([header] + [row] * 1001).encode()))
        self.assertIn('2 MB', read('a.csv', b'x' * (spreadsheet.MAX_FILE_BYTES + 1)))
        self.assertIn('not an .xlsx', read('a.xlsx', b'not a zip'))
        self.assertIn('.csv or .xlsx', read('a.xls', b'x'))

        bomb = io.BytesIO()
        with zipfile.ZipFile(bomb, 'w', zipfile.ZIP_DEFLATED) as archive:
            archive.writestr('xl/sharedStrings.xml', b'0' * (spreadsheet.MAX_UNZIPPED_BYTES + 1))
        self.assertIn('too large', read('a.xlsx', bomb.getvalue()))

        workbook = openpyxl.Workbook()
        workbook.active.append(COLUMNS)
        workbook.active.append(['', '=1+1'])
        out = io.BytesIO()
        workbook.save(out)
        self.assertIn('formulas', read('a.xlsx', out.getvalue()))

    #
    # Atomic confirm, history and idempotency
    #
    def test_failed_confirm_rolls_back_everything(self):
        rows = synthetic_rows()
        rows[-1]['notes'] = 'boom'
        upload = csv_file(rows)
        plan_hash = self.post(PREVIEW, upload, mode='create').data['plan_hash']
        before = (plan_counts(), history())
        save = SlotEntry.save

        def failing_save(entry, *args, **kwargs):
            if entry.comment == 'boom':
                raise RuntimeError('boom')
            return save(entry, *args, **kwargs)

        with patch.object(SlotEntry, 'save', failing_save), self.assertRaises(RuntimeError):
            self.post(CONFIRM, upload, mode='create', plan_hash=plan_hash)
        self.assertEqual((plan_counts(), history()), before)

    def test_deletes_never_touch_history(self):
        routine = self.import_plan(csv_file(synthetic_rows()))
        upper, rest, lower = routine.days.order_by('order')
        logged = upper.slots.get(order=1).entries.get()
        log = WorkoutLog.objects.filter(user__username='admin').first()
        session = WorkoutSession.objects.filter(user__username='admin').exclude(logs=log).first()
        WorkoutLog.objects.filter(pk=log.pk).update(slot_entry=logged, routine=routine)
        WorkoutSession.objects.filter(pk=session.pk).update(day=lower, routine=routine)
        before = history()
        rows = download_rows(self.export(routine))

        # Dropping a logged entry, a day with a session, or making the logged day rest
        for changed in (
            rows[1:],
            [r for r in rows if r['day_order'] != '3'],
            [
                {
                    **rows[0],
                    **{c: '' for c in spreadsheet.ENTRY_COLUMNS},
                    'slot_id': '',
                    'slot_order': '',
                    'exercise_id': '',
                    'exercise_uuid': '',
                    'exercise_name': '',
                    'day_is_rest': 'yes',
                }
            ]
            + rows[4:],
        ):
            upload = csv_file(changed)
            data = self.post(PREVIEW, upload, mode='update', routine=routine.pk).data
            self.assertFalse(data['ok'])
            self.assertTrue(data['diff']['blocked_deletes'])
            confirm = self.post(
                CONFIRM, upload, mode='update', routine=routine.pk, plan_hash=data['plan_hash']
            )
            self.assertEqual(confirm.status_code, status.HTTP_400_BAD_REQUEST)
            self.assertEqual(history(), before)

        # An entry without logs can go
        unlogged = [r for r in rows if r['exercise_name'] != 'Pending exercise']
        self.import_plan(csv_file(unlogged), 'update', status.HTTP_200_OK, routine=routine.pk)
        self.assertFalse(upper.slots.filter(order=3).exists())
        self.assertEqual(history(), before)

    def test_stale_or_repeated_confirm_is_refused(self):
        routine = self.import_plan(csv_file(synthetic_rows()))
        rows = download_rows(self.export(routine))
        rows[0]['day_name'] = 'Push'
        for row in rows[1:4]:
            row['day_name'] = 'Push'
        upload = csv_file(rows)
        preview = self.post(PREVIEW, upload, mode='update', routine=routine.pk).data

        # The routine changed after the preview
        Day.objects.filter(routine=routine, order=3).update(description='edited')
        stale = self.post(
            CONFIRM, upload, mode='update', routine=routine.pk, plan_hash=preview['plan_hash']
        )
        self.assertEqual(stale.status_code, status.HTTP_409_CONFLICT)
        self.assertFalse(routine.days.filter(name='Push').exists())

        # Applying the same preview twice
        preview = self.post(PREVIEW, upload, mode='update', routine=routine.pk).data
        args = {'mode': 'update', 'routine': routine.pk, 'plan_hash': preview['plan_hash']}
        self.assertEqual(self.post(CONFIRM, upload, **args).status_code, status.HTTP_200_OK)
        self.assertEqual(self.post(CONFIRM, upload, **args).status_code, status.HTTP_409_CONFLICT)

        # Creating the same file twice
        again = self.post(PREVIEW, csv_file(synthetic_rows()), mode='create').data
        self.assertIn((2, 'routine_name'), {(e['row'], e['column']) for e in again['errors']})

    #
    # Access
    #
    def test_access(self):
        mine = csv_file(synthetic_rows())
        # Update of someone else's routine, a public template or a template of mine
        for pk in (2, 5):
            response = self.post(PREVIEW, mine, mode='update', routine=pk)
            self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)
            response = self.post(CONFIRM, mine, mode='update', routine=pk, plan_hash='x')
            self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)
        Routine.objects.filter(pk=1).update(is_template=True)
        self.assertEqual(
            self.post(PREVIEW, mine, mode='update', routine=1).status_code,
            status.HTTP_404_NOT_FOUND,
        )

        # Export follows the structure endpoint: owner, or read access to public templates
        self.assertEqual(self.client.get('/api/v2/routine/2/export/').status_code, 404)
        self.assertEqual(self.client.get('/api/v2/routine/5/export/?file=xlsx').status_code, 200)
        self.assertEqual(self.client.get('/api/v2/routine/1/export/?file=pdf').status_code, 400)

        self.client.force_authenticate(user=None)
        self.assertEqual(self.client.get('/api/v2/routine/1/export/').status_code, 403)
        self.assertEqual(self.post(PREVIEW, mine, mode='create').status_code, 403)
        self.assertEqual(self.client.get('/api/v2/routine/import-template/').status_code, 403)

    def test_template(self):
        response = self.client.get('/api/v2/routine/import-template/?file=csv')
        self.assertEqual(
            response['Content-Disposition'], 'attachment; filename="routine-template.csv"'
        )
        self.assertEqual(response.content.decode('utf-8-sig').strip(), ','.join(COLUMNS))

        response = self.client.get('/api/v2/routine/import-template/?file=xlsx')
        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, ['routine', 'help'])
        self.assertEqual([c.value for c in workbook['routine'][1]], COLUMNS)
