# This file is part of wger Workout Manager.
#
# wger Workout Manager is free software: you can redistribute it and/or modify
# it under the terms of the GNU Affero General Public License as published by
# the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version.
#
# wger Workout Manager is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU General Public License for more details.
#
# You should have received a copy of the GNU Affero General Public License
# along with Workout Manager.  If not, see <http://www.gnu.org/licenses/>.

"""
CSV and XLSX import and export of the planned structure of a routine

This handles plans only. Workout sessions and logs are never written here, they
are only read to refuse any delete that would cascade into logged history.
"""

# Standard Library
import csv
import datetime
import hashlib
import io
import json
import uuid
import zipfile
from dataclasses import (
    dataclass,
    field,
)
from decimal import (
    Decimal,
    InvalidOperation,
)

# Django
from django.db import (
    DEFAULT_DB_ALIAS,
    transaction,
)
from django.db.models import (
    ProtectedError,
    RestrictedError,
)
from django.db.models.deletion import Collector
from django.http import HttpResponse
from django.utils import timezone

# Third Party
import openpyxl

# wger
from wger.core.models import (
    RepetitionUnit,
    WeightUnit,
)
from wger.exercises.models import (
    Alias,
    Exercise,
    Translation,
)
from wger.manager.api.serializers import (
    RoutineSerializer,
    RoutineStructureSerializer,
)
from wger.manager.consts import (
    REP_UNIT_REPETITIONS,
    RIR_OPTIONS,
    WEIGHT_UNIT_KG,
)
from wger.manager.helpers import reset_routine_cache
from wger.manager.models import (
    Day,
    MaxRepetitionsConfig,
    MaxRestConfig,
    MaxRiRConfig,
    MaxSetsConfig,
    MaxWeightConfig,
    RepetitionsConfig,
    RestConfig,
    RiRConfig,
    Routine,
    SetsConfig,
    Slot,
    SlotEntry,
    WeightConfig,
    WorkoutSessionRecovery,
)
from wger.manager.models.slot_entry import ExerciseType
from wger.utils.constants import ENGLISH_SHORT_NAME
from wger.utils.language import load_language


COLUMNS = [
    'routine_id',
    'routine_name',
    'routine_description',
    'start',
    'end',
    'day_id',
    'day_order',
    'day_name',
    'day_description',
    'day_is_rest',
    'slot_id',
    'slot_order',
    'slot_comment',
    'entry_id',
    'entry_order',
    'exercise_id',
    'exercise_uuid',
    'exercise_name',
    'set_type',
    'sets',
    'reps',
    'max_reps',
    'rep_unit',
    'weight',
    'max_weight',
    'weight_unit',
    'rir',
    'max_rir',
    'rest',
    'max_rest',
    'notes',
    'unsupported',
]
ROUTINE_COLUMNS = ('routine_name', 'routine_description', 'start', 'end')
SLOT_COLUMNS = ('slot_id', 'slot_order', 'slot_comment')
EXERCISE_COLUMNS = ('exercise_id', 'exercise_uuid', 'exercise_name')

CONFIG_COLUMNS = {
    'sets': SetsConfig,
    'reps': RepetitionsConfig,
    'max_reps': MaxRepetitionsConfig,
    'weight': WeightConfig,
    'max_weight': MaxWeightConfig,
    'rir': RiRConfig,
    'max_rir': MaxRiRConfig,
    'rest': RestConfig,
    'max_rest': MaxRestConfig,
}
"""Columns stored as the iteration 1 row of a config table"""

ENTRY_COLUMNS = ('entry_id', 'entry_order', 'set_type', 'rep_unit', 'weight_unit', 'notes')
ENTRY_COLUMNS += tuple(CONFIG_COLUMNS)

CONFIG_SETS = [f'{model.__name__.lower()}_set' for model in CONFIG_COLUMNS.values()]
CONFIG_SETS.append('maxsetsconfig_set')

BOUNDS = {
    'sets': (1, 50, True),
    'reps': (0, 3000, False),
    'max_reps': (0, 3000, False),
    'weight': (0, 3000, False),
    'max_weight': (0, 3000, False),
    'rest': (0, 1800, True),
    'max_rest': (0, 600, True),
}
"""(min, max, whole number) mirroring the config model validators"""

RANGES = (('reps', 'max_reps'), ('weight', 'max_weight'), ('rir', 'max_rir'), ('rest', 'max_rest'))

PLAN_MODELS = {Day, Slot, SlotEntry, MaxSetsConfig, *CONFIG_COLUMNS.values()}
"""Everything a plan delete may remove, anything else is history"""

MAX_FILE_BYTES = 2 * 1024 * 1024
MAX_UNZIPPED_BYTES = 20 * 1024 * 1024
MAX_ZIP_MEMBERS = 200
MAX_ROWS = 1000
MAX_COLUMNS = 64
MAX_SCANNED_ROWS = 5000
ACTIVE_PARTS = ('vbaproject', 'activex', 'externallink', 'embeddings', 'macrosheets')

FORMULA_START = ('=', '+', '-', '@')
ESCAPED_START = FORMULA_START + ('\t', '\r', "'")

CSV_TYPE = 'text/csv; charset=utf-8'
XLSX_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


class SpreadsheetError(Exception):
    """The file as a whole can't be read"""


#
# Cells
#
def escape(value: str) -> str:
    """Prefix text a spreadsheet would run as a formula (or an escape itself)"""
    return f"'{value}" if value[:1] in ESCAPED_START else value


def unescape(value: str) -> str:
    """Inverse of `escape`, unescaped formulas are refused"""
    value = value.strip()
    if value[:1] in FORMULA_START:
        raise ValueError("Starts like a formula; prefix it with ' to keep it as text.")
    if value[:1] == "'" and value[1:2] in ESCAPED_START:
        return value[1:]
    return value


def _fmt(value: Decimal) -> str:
    return format(value.normalize(), 'f')


def _text(value) -> str:
    """A cell of an uploaded workbook as text"""
    if value is None:
        return ''
    if isinstance(value, bool):
        return 'yes' if value else 'no'
    if isinstance(value, float):
        return _fmt(Decimal(repr(value)))
    if isinstance(value, datetime.datetime):
        return value.date().isoformat()
    if isinstance(value, datetime.date):
        return value.isoformat()
    return str(value)


def _int(text: str, low: int = 1) -> int:
    try:
        value = int(text)
    except ValueError:
        raise ValueError('Must be a whole number.')
    if value < low:
        raise ValueError(f'Must be at least {low}.')
    return value


def _number(column: str, text: str) -> Decimal:
    try:
        value = Decimal(text)
    except InvalidOperation:
        raise ValueError('Must be a number.')
    if not value.is_finite():
        raise ValueError('Must be a number.')

    if column in ('rir', 'max_rir'):
        if value not in RIR_OPTIONS[1:]:
            raise ValueError(f'Must be one of {", ".join(str(o) for o in RIR_OPTIONS[1:])}.')
        return value

    low, high, whole = BOUNDS[column]
    if whole and value != value.to_integral_value():
        raise ValueError('Must be a whole number.')
    if value.as_tuple().exponent < -2:
        raise ValueError('At most 2 decimals.')
    if not low <= value <= high:
        raise ValueError(f'Must be between {low} and {high}.')
    return value


def _yes_no(text: str) -> bool:
    value = text.lower()
    if value in ('', 'no', 'false'):
        return False
    if value in ('yes', 'true'):
        return True
    raise ValueError('Must be yes or no.')


#
# Reading
#
def _read_csv(data: bytes) -> list[list[str]]:
    try:
        return list(csv.reader(io.StringIO(data.decode('utf-8-sig'), newline='')))
    except (UnicodeDecodeError, csv.Error):
        raise SpreadsheetError('The file is not a UTF-8 CSV file.')


def _read_xlsx(data: bytes) -> list[list[str]]:
    buffer = io.BytesIO(data)
    if not zipfile.is_zipfile(buffer):
        raise SpreadsheetError('The file is not an .xlsx workbook.')
    try:
        with zipfile.ZipFile(buffer) as archive:
            members = archive.infolist()
    except zipfile.BadZipFile:
        raise SpreadsheetError('The file is not an .xlsx workbook.')

    # Checked before openpyxl inflates anything
    if len(members) > MAX_ZIP_MEMBERS or sum(m.file_size for m in members) > MAX_UNZIPPED_BYTES:
        raise SpreadsheetError('The workbook is too large once unpacked.')
    if any(part in m.filename.lower() for m in members for part in ACTIVE_PARTS):
        raise SpreadsheetError('Macros, external links and embedded objects are not accepted.')

    try:
        workbook = openpyxl.load_workbook(buffer, read_only=True, data_only=False)
    except Exception:
        raise SpreadsheetError('The workbook could not be read.')

    grid = []
    try:
        sheet = workbook.worksheets[0]
        # Ignore the declared dimensions, they can claim any size
        sheet.reset_dimensions()
        for number, cells in enumerate(sheet.iter_rows(max_col=MAX_COLUMNS), start=1):
            if number > MAX_SCANNED_ROWS:
                raise SpreadsheetError(f'At most {MAX_ROWS} rows can be imported.')
            for cell in cells:
                if cell.data_type in ('f', 'e') or not isinstance(
                    cell.value, (str, int, float, bool, datetime.date, type(None))
                ):
                    raise SpreadsheetError(
                        f'Row {number}: formulas and error values are not accepted, '
                        'enter plain values.'
                    )
            grid.append([_text(cell.value) for cell in cells])
    finally:
        workbook.close()
    return grid


def read_file(upload) -> tuple[list[dict], list[dict]]:
    """
    Read an uploaded CSV or XLSX template into rows keyed by column

    Returns the rows and the per cell errors. Raises SpreadsheetError if the
    file as a whole can't be used.
    """
    if upload is None:
        raise SpreadsheetError('No file was uploaded.')
    if upload.size > MAX_FILE_BYTES:
        raise SpreadsheetError('The file is larger than 2 MB.')

    name = (upload.name or '').lower()
    data = upload.read(MAX_FILE_BYTES + 1)
    if name.endswith('.xlsx'):
        grid = _read_xlsx(data)
    elif name.endswith('.csv'):
        grid = _read_csv(data)
    else:
        raise SpreadsheetError('Upload a .csv or .xlsx file.')

    if not grid:
        raise SpreadsheetError('The file is empty.')
    header = [value.strip().lower() for value in grid[0]]
    while header and not header[-1]:
        header.pop()
    unknown = [h or '(blank)' for h in header if h not in COLUMNS]
    missing = [c for c in COLUMNS if c not in header]
    duplicate = sorted({h for h in header if header.count(h) > 1})
    if unknown or missing or duplicate:
        problems = [
            f'{label}: {", ".join(names)}'
            for label, names in (
                ('unknown columns', unknown),
                ('missing columns', missing),
                ('duplicate columns', duplicate),
            )
            if names
        ]
        raise SpreadsheetError(
            f'The first row must be the template header ({"; ".join(problems)}).'
        )

    rows, errors = [], []
    for number, values in enumerate(grid[1:], start=2):
        if not any(value.strip() for value in values):
            continue
        if number > MAX_ROWS + 1:
            raise SpreadsheetError(f'At most {MAX_ROWS} rows can be imported.')
        if any(value.strip() for value in values[len(header) :]):
            errors.append({'row': number, 'column': None, 'message': 'Values outside the columns.'})

        row = {'_row': number}
        values = values + [''] * (len(header) - len(values))
        for column, value in zip(header, values):
            try:
                row[column] = unescape(value)
            except ValueError as e:
                errors.append({'row': number, 'column': column, 'message': str(e)})
                row[column] = ''
        rows.append(row)
    return rows, errors


#
# Planning
#
@dataclass
class Plan:
    mode: str
    target: Routine | None
    drop_unsupported: bool = False
    errors: list = field(default_factory=list)
    rows: list = field(default_factory=list)
    routine: dict | None = None
    days: list = field(default_factory=list)
    deletes: dict = field(default_factory=lambda: {'days': [], 'slots': [], 'entries': []})
    blocked: list = field(default_factory=list)
    dropped: list = field(default_factory=list)
    hash: str = ''

    @property
    def ok(self) -> bool:
        return not self.errors and not self.blocked and bool(self.days) and self.routine is not None

    def preview(self) -> dict:
        def count(key):
            return {
                'days': sum(1 for d in self.days if (d['id'] is None) == key),
                'slots': sum(1 for d in self.days for s in d['slots'] if (s['id'] is None) == key),
                'entries': sum(
                    1
                    for d in self.days
                    for s in d['slots']
                    for e in s['entries']
                    if (e['id'] is None) == key
                ),
            }

        return {
            'ok': self.ok,
            'plan_hash': self.hash,
            'mode': self.mode,
            'routine': self.routine
            and {'id': self.target.pk if self.target else None, **self.routine},
            'rows': self.rows,
            'errors': self.errors,
            'diff': {
                'create': count(True),
                'update': count(False),
                'delete': {key: len(ids) for key, ids in self.deletes.items()},
                'blocked_deletes': self.blocked,
            },
            'unsupported_dropped': self.dropped,
        }


def _exercise_names(exercise_ids, language: str) -> dict[int, str]:
    names = {}
    for exercise in Exercise.objects.filter(pk__in=exercise_ids).prefetch_related(
        'translations__language'
    ):
        translations = list(exercise.translations.all())
        match = [t for t in translations if t.language.short_name == language] or [
            t for t in translations if t.language.short_name == ENGLISH_SHORT_NAME
        ]
        best = (match or translations or [None])[0]
        names[exercise.pk] = best.name if best else ''
    return names


def _resolve_exercise(row: dict, languages: list[str]) -> tuple[Exercise | None, dict, tuple]:
    """
    Returns the exercise, the preview info and an optional (column, message) error

    Never guesses: ids and uuids must exist, names must match exactly one
    exercise, and a name next to an id must belong to it.
    """
    ident, ident_uuid, name = (row[c] for c in EXERCISE_COLUMNS)
    exercise = None
    if ident or ident_uuid:
        column = 'exercise_id' if ident else 'exercise_uuid'
        try:
            lookup = {'pk': int(ident)} if ident else {'uuid': uuid.UUID(ident_uuid)}
        except ValueError:
            return None, {'how': 'unresolved'}, (column, 'Not a valid id.')
        exercise = Exercise.objects.filter(**lookup).first()
        if exercise is None:
            return None, {'how': 'unresolved'}, (column, 'No exercise with this id.')
        how = 'id' if ident else 'uuid'
        if name and not (
            exercise.translations.filter(name__iexact=name).exists()
            or Alias.objects.filter(translation__exercise=exercise, alias__iexact=name).exists()
        ):
            return (
                None,
                {'how': 'mismatch', 'id': exercise.pk},
                (
                    'exercise_name',
                    'Name/id mismatch; clear exercise_id and exercise_uuid to rematch.',
                ),
            )
    else:
        ids = set()
        for language in languages:
            ids = set(
                Translation.objects.filter(
                    name__iexact=name, language__short_name=language
                ).values_list('exercise_id', flat=True)
            ) | set(
                Alias.objects.filter(
                    alias__iexact=name, translation__language__short_name=language
                ).values_list('translation__exercise_id', flat=True)
            )
            if ids:
                break
        if not ids:
            return None, {'how': 'unresolved'}, ('exercise_name', 'No exercise with this name.')
        if len(ids) > 1:
            names = _exercise_names(sorted(ids)[:5], languages[0])
            candidates = [{'id': pk, 'name': n} for pk, n in names.items()]
            return (
                None,
                {'how': 'ambiguous', 'candidates': candidates},
                ('exercise_name', 'Several exercises have this name; set exercise_id.'),
            )
        exercise = Exercise.objects.get(pk=ids.pop())
        how = 'name'

    info = {'id': exercise.pk, 'name': _exercise_names([exercise.pk], languages[0])[exercise.pk]}
    return exercise, {**info, 'how': how}, ()


def _recovery_references(user) -> tuple[set, set]:
    """Day and slot entry ids a restorable deleted workout still points to"""
    days, entries = set(), set()
    for snapshot in WorkoutSessionRecovery.objects.filter(
        user=user, expires_at__gt=timezone.now()
    ).values_list('snapshot', flat=True):
        days.add(str((snapshot.get('session') or {}).get('day_id')))
        entries.update(str(log.get('slot_entry_id')) for log in snapshot.get('logs') or [])
    return days, entries


def _history_reason(obj, recovery: tuple[set, set]) -> str | None:
    """Why deleting this plan object would touch history, None if it wouldn't"""
    collector = Collector(using=DEFAULT_DB_ALIAS)
    try:
        collector.collect([obj])
    except (ProtectedError, RestrictedError):
        return 'Other records depend on it.'

    touched = {model for model, objs in collector.data.items() if objs}
    touched |= {qs.model for qs in collector.fast_deletes if qs.exists()}
    touched |= {
        f.model
        for (f, _), groups in collector.field_updates.items()
        if any(g.exists() if hasattr(g, 'exists') else g for g in groups)
    }
    history = touched - PLAN_MODELS
    if history:
        names = ', '.join(sorted(str(model._meta.verbose_name_plural) for model in history))
        return f'Deleting it would delete logged history ({names}).'

    days = {str(d.pk) for d in collector.data.get(Day, ())}
    entries = {str(e.pk) for e in collector.data.get(SlotEntry, ())}
    if days & recovery[0] or entries & recovery[1]:
        return 'A deleted workout that can still be restored refers to it.'
    return None


def build_plan(upload, user, mode: str, target: Routine | None = None, drop_unsupported=False):
    """
    Validate an upload against the current data and resolve it into a plan

    `mode` is create or update, `target` the owned routine for update.
    Nothing is written.
    """
    plan = Plan(mode=mode, target=target, drop_unsupported=drop_unsupported)
    try:
        rows, plan.errors = read_file(upload)
    except SpreadsheetError as e:
        plan.errors.append({'row': None, 'column': None, 'message': str(e)})
        return plan
    if not rows:
        plan.errors.append({'row': None, 'column': None, 'message': 'The file has no rows.'})
        return plan

    def err(row, column, message):
        plan.errors.append({'row': row, 'column': column, 'message': message})

    def parse_id(row, column):
        if not row[column]:
            return None
        if mode == 'create':
            err(row['_row'], column, 'Must be blank when creating; use update for this routine.')
            return None
        try:
            return _int(row[column])
        except ValueError as e:
            err(row['_row'], column, str(e))
            return None

    languages = list(dict.fromkeys([load_language().short_name, ENGLISH_SHORT_NAME]))
    rep_units = {u.name.lower(): u.pk for u in RepetitionUnit.objects.all()}
    weight_units = {u.name.lower(): u.pk for u in WeightUnit.objects.all()}

    existing_days, existing_slots, existing_entries = {}, {}, {}
    if target is not None:
        prefetch = ['slots__entries', *[f'slots__entries__{name}' for name in CONFIG_SETS]]
        existing_days = {d.pk: d for d in target.days.prefetch_related(*prefetch)}
        existing_slots = {s.pk: s for d in existing_days.values() for s in d.slots.all()}
        existing_entries = {e.pk: e for s in existing_slots.values() for e in s.entries.all()}

    first = rows[0]
    days = {}
    seen_ids = {'day_id': set(), 'slot_id': set(), 'entry_id': set()}
    exercise_cache = {}

    def claim(row, column, value):
        """One id may only stand for one group"""
        if value is None:
            return
        if value in seen_ids[column]:
            err(row['_row'], column, 'This id is used by another group.')
        seen_ids[column].add(value)

    def group(existing, values, row, name):
        for key, value in values.items():
            if existing[key] != value:
                err(row['_row'], key, f'Conflicts with row {existing["row"]} for this {name}.')

    for row in rows:
        number = row['_row']
        info = None
        for column in ROUTINE_COLUMNS:
            if row[column] and row[column] != first[column]:
                err(number, column, 'Must be blank or match the first row.')
        if row['routine_id']:
            if mode == 'create':
                err(
                    number,
                    'routine_id',
                    'Must be blank when creating; use update for this routine.',
                )
            elif row['routine_id'] != str(target.pk):
                err(number, 'routine_id', 'Does not match the routine being updated.')

        if mode == 'create' and row['unsupported']:
            codes = [code for code in row['unsupported'].split(';') if code.strip()]
            if drop_unsupported:
                plan.dropped.append({'row': number, 'codes': codes})
            else:
                err(
                    number,
                    'unsupported',
                    f'This row has settings the file cannot carry ({", ".join(codes)}). '
                    'Import with drop_unsupported to leave them out.',
                )

        try:
            day_order = _int(row['day_order'])
        except ValueError as e:
            err(number, 'day_order', str(e))
            plan.rows.append({'row': number, 'exercise': None})
            continue
        try:
            is_rest = _yes_no(row['day_is_rest'])
        except ValueError as e:
            err(number, 'day_is_rest', str(e))
            plan.rows.append({'row': number, 'exercise': None})
            continue

        day_values = {
            'day_id': parse_id(row, 'day_id'),
            'day_name': row['day_name'],
            'day_description': row['day_description'],
            'day_is_rest': is_rest,
        }
        day = days.get(day_order)
        if day is None:
            day = days[day_order] = {**day_values, 'row': number, 'slots': {}, 'bare': None}
            claim(row, 'day_id', day['day_id'])
            if day['day_id'] is not None and day['day_id'] not in existing_days:
                err(number, 'day_id', 'Not a day of this routine.')
            for column, model_field in (('day_name', 'name'), ('day_description', 'description')):
                limit = Day._meta.get_field(model_field).max_length
                if len(row[column]) > limit:
                    err(number, column, f'At most {limit} characters.')
        else:
            group(day, day_values, row, 'day')
            if is_rest:
                err(number, 'day_order', 'A rest day has exactly one row.')

        has_slot = any(row[c] for c in SLOT_COLUMNS)
        has_exercise = any(row[c] for c in EXERCISE_COLUMNS)
        has_entry = any(row[c] for c in ENTRY_COLUMNS)

        if is_rest:
            if has_slot or has_exercise or has_entry:
                err(number, 'day_is_rest', 'A rest day row cannot have slots or exercises.')
        elif not (has_slot or has_exercise or has_entry):
            if day['bare'] is not None or day['slots']:
                err(number, 'slot_order', 'Remove this empty row, the day has other rows.')
            day['bare'] = number
        elif not row['slot_order']:
            err(number, 'slot_order', 'Required for exercises.')
        elif day['bare'] is not None:
            err(day['bare'], 'slot_order', 'Remove this empty row, the day has other rows.')
            day['bare'] = None
        if is_rest or not row['slot_order'] or not (has_slot or has_exercise or has_entry):
            plan.rows.append({'row': number, 'exercise': None})
            continue

        try:
            slot_order = _int(row['slot_order'])
        except ValueError as e:
            err(number, 'slot_order', str(e))
            plan.rows.append({'row': number, 'exercise': None})
            continue
        slot_values = {'slot_id': parse_id(row, 'slot_id'), 'slot_comment': row['slot_comment']}
        slot = day['slots'].get(slot_order)
        if slot is None:
            slot = day['slots'][slot_order] = {
                **slot_values,
                'row': number,
                'entries': {},
                'bare': None,
            }
            claim(row, 'slot_id', slot['slot_id'])
            if slot['slot_id'] is not None and (
                slot['slot_id'] not in existing_slots
                or existing_slots[slot['slot_id']].day_id != day['day_id']
            ):
                err(number, 'slot_id', 'Not a slot of this day.')
            limit = Slot._meta.get_field('comment').max_length
            if len(row['slot_comment']) > limit:
                err(number, 'slot_comment', f'At most {limit} characters.')
        else:
            group(slot, slot_values, row, 'slot')

        if not (has_exercise or has_entry):
            if slot['bare'] is not None or slot['entries']:
                err(number, 'exercise_name', 'Remove this empty row, the slot has other rows.')
            slot['bare'] = number
            plan.rows.append({'row': number, 'exercise': None})
            continue
        if slot['bare'] is not None:
            err(slot['bare'], 'exercise_name', 'Remove this empty row, the slot has other rows.')
            slot['bare'] = None
        if not has_exercise:
            err(number, 'exercise_name', 'An exercise is required.')
            plan.rows.append({'row': number, 'exercise': None})
            continue

        # Exercise entry
        key = tuple(row[c].lower() for c in EXERCISE_COLUMNS)
        if key not in exercise_cache:
            exercise_cache[key] = _resolve_exercise(row, languages)
        exercise, info, problem = exercise_cache[key]
        if problem:
            err(number, *problem)

        entry_id = parse_id(row, 'entry_id')
        claim(row, 'entry_id', entry_id)
        existing_entry = None
        if entry_id is not None:
            existing_entry = existing_entries.get(entry_id)
            if existing_entry is None or existing_entry.slot_id != slot['slot_id']:
                err(number, 'entry_id', 'Not an exercise of this slot.')
                existing_entry = None

        entry = {'id': entry_id, 'exercise': exercise.pk if exercise else None, 'configs': {}}
        try:
            entry['order'] = _int(row['entry_order'])
            if entry['order'] in slot['entries']:
                err(number, 'entry_order', 'Used twice in this slot.')
        except ValueError as e:
            err(number, 'entry_order', 'Required.' if not row['entry_order'] else str(e))
            entry['order'] = None

        entry['type'] = row['set_type'].lower() or ExerciseType.NORMAL
        if entry['type'] not in ExerciseType.values:
            err(number, 'set_type', f'Must be one of {", ".join(ExerciseType.values)}.')
        for column, units, default in (
            ('rep_unit', rep_units, REP_UNIT_REPETITIONS),
            ('weight_unit', weight_units, WEIGHT_UNIT_KG),
        ):
            entry[column] = units.get(row[column].lower(), None) if row[column] else default
            if entry[column] is None:
                err(number, column, f'Unknown unit, use one of: {", ".join(sorted(units))}.')
        entry['comment'] = row['notes']
        limit = SlotEntry._meta.get_field('comment').max_length
        if len(row['notes']) > limit:
            err(number, 'notes', f'At most {limit} characters.')

        values = {}
        for column, model in CONFIG_COLUMNS.items():
            value = None
            if row[column]:
                try:
                    value = _number(column, row[column])
                except ValueError as e:
                    err(number, column, str(e))
            elif column == 'sets':
                err(number, column, 'Required.')
            elif existing_entry is not None:
                configs = list(getattr(existing_entry, f'{model.__name__.lower()}_set').all())
                if any(c.iteration == 1 for c in configs) and any(
                    c.iteration > 1 or c.requirements for c in configs
                ):
                    err(number, column, 'Cannot be cleared, it has a progression or requirements.')
            values[column] = value
            entry['configs'][column] = None if value is None else _fmt(value)
        for low, high in RANGES:
            if values[low] is not None and values[high] is not None and values[high] < values[low]:
                err(number, high, f'Must not be below {low}.')

        if entry['order'] is not None and entry['order'] not in slot['entries']:
            slot['entries'][entry['order']] = entry
        plan.rows.append({'row': number, 'exercise': info})

    # Routine
    if not first['routine_name']:
        err(first['_row'], 'routine_name', 'Required on the first row.')
    serializer = RoutineSerializer(
        instance=target,
        data={
            'name': first['routine_name'],
            'description': first['routine_description'],
            'start': first['start'] or None,
            'end': first['end'] or None,
        },
    )
    if serializer.is_valid():
        data = serializer.validated_data
        plan.routine = {
            'name': data['name'],
            'description': data.get('description', ''),
            'start': data['start'].isoformat(),
            'end': data['end'].isoformat(),
        }
        if (
            mode == 'create'
            and Routine.objects.filter(user=user, name__iexact=data['name']).exists()
        ):
            err(first['_row'], 'routine_name', 'You already have a routine with this name.')
    else:
        columns = {'name': 'routine_name', 'description': 'routine_description'}
        for key, messages in serializer.errors.items():
            for message in messages:
                err(
                    first['_row'],
                    columns.get(key, key if key in ('start', 'end') else None),
                    message,
                )

    plan.days = [
        {
            'id': day['day_id'],
            'order': order,
            'name': day['day_name'],
            'description': day['day_description'],
            'is_rest': day['day_is_rest'],
            'slots': [
                {
                    'id': slot['slot_id'],
                    'order': slot_order,
                    'comment': slot['slot_comment'],
                    'entries': [
                        {
                            'id': entry['id'],
                            'order': entry_order,
                            'exercise': entry['exercise'],
                            'type': entry['type'],
                            'repetition_unit': entry['rep_unit'],
                            'weight_unit': entry['weight_unit'],
                            'comment': entry['comment'],
                            'configs': entry['configs'],
                        }
                        for entry_order, entry in sorted(slot['entries'].items())
                    ],
                }
                for slot_order, slot in sorted(day['slots'].items())
            ],
        }
        for order, day in sorted(days.items())
    ]

    # Everything of the target the file no longer lists is deleted, never history
    if target is not None:
        kept = {
            'days': {d['id'] for d in plan.days},
            'slots': {s['id'] for d in plan.days for s in d['slots']},
            'entries': {e['id'] for d in plan.days for s in d['slots'] for e in s['entries']},
        }
        for day in existing_days.values():
            if day.pk not in kept['days']:
                plan.deletes['days'].append(day)
                continue
            for slot in day.slots.all():
                if slot.pk not in kept['slots']:
                    plan.deletes['slots'].append(slot)
                    continue
                for entry in slot.entries.all():
                    if entry.pk not in kept['entries']:
                        plan.deletes['entries'].append(entry)

        recovery = _recovery_references(user)
        for kind, objs in plan.deletes.items():
            for obj in objs:
                reason = _history_reason(obj, recovery)
                if reason:
                    plan.blocked.append({'kind': kind[:-1], 'id': obj.pk, 'reason': reason})
        plan.deletes = {kind: [obj.pk for obj in objs] for kind, objs in plan.deletes.items()}

    error_rows = {e['row'] for e in plan.errors}
    for row in plan.rows:
        row['status'] = 'error' if row['row'] in error_rows else 'ok'

    if target is not None:
        version = json.dumps(RoutineStructureSerializer(target).data, sort_keys=True, default=str)
    else:
        version = sorted(
            n.lower() for n in Routine.objects.filter(user=user).values_list('name', flat=True)
        )
    resolved = {
        'mode': mode,
        'target': target.pk if target else None,
        'drop_unsupported': drop_unsupported,
        'routine': plan.routine,
        'days': plan.days,
        'deletes': plan.deletes,
        'version': version,
    }
    plan.hash = hashlib.sha256(
        json.dumps(resolved, sort_keys=True, default=str).encode()
    ).hexdigest()
    return plan


def apply_plan(plan: Plan, user) -> Routine:
    """
    Write a checked plan. The caller holds the owner lock inside the
    transaction and rebuilt the plan under it.
    """
    assert plan.ok

    routine = plan.target
    if routine is None:
        routine = Routine(user=user, is_template=False)
    for key, value in plan.routine.items():
        setattr(
            routine, key, datetime.date.fromisoformat(value) if key in ('start', 'end') else value
        )
    routine.save()

    # Guarded against history in build_plan, deepest first
    SlotEntry.objects.filter(pk__in=plan.deletes['entries'], slot__day__routine=routine).delete()
    Slot.objects.filter(pk__in=plan.deletes['slots'], day__routine=routine).delete()
    Day.objects.filter(pk__in=plan.deletes['days'], routine=routine).delete()

    for day_data in plan.days:
        day = (
            Day.objects.get(pk=day_data['id'], routine=routine)
            if day_data['id']
            else Day(routine=routine)
        )
        day.order = day_data['order']
        day.name = day_data['name']
        day.description = day_data['description']
        day.is_rest = day_data['is_rest']
        day.save()

        for slot_data in day_data['slots']:
            slot = (
                Slot.objects.get(pk=slot_data['id'], day=day) if slot_data['id'] else Slot(day=day)
            )
            slot.order = slot_data['order']
            slot.comment = slot_data['comment']
            slot.save()

            for entry_data in slot_data['entries']:
                entry = (
                    SlotEntry.objects.get(pk=entry_data['id'], slot=slot)
                    if entry_data['id']
                    else SlotEntry(slot=slot)
                )
                entry.order = entry_data['order']
                entry.exercise_id = entry_data['exercise']
                entry.type = entry_data['type']
                entry.repetition_unit_id = entry_data['repetition_unit']
                entry.weight_unit_id = entry_data['weight_unit']
                entry.comment = entry_data['comment']
                entry.save()

                # Only the base value, progressions and requirements stay as they are
                for column, value in entry_data['configs'].items():
                    model = CONFIG_COLUMNS[column]
                    if value is None:
                        model.objects.filter(slot_entry=entry, iteration=1).delete()
                    else:
                        model.objects.update_or_create(
                            slot_entry=entry, iteration=1, defaults={'value': Decimal(value)}
                        )

    transaction.on_commit(lambda: reset_routine_cache(routine))
    return routine


#
# Writing
#
def _codes(obj, *extra) -> list[str]:
    return [code for code, flag in (('json_config', obj.config is not None), *extra) if flag]


def export_rows(routine: Routine) -> list[list]:
    """One row per exercise entry, empty days and slots and rest days get one row"""
    language = load_language().short_name
    profile = routine.user.userprofile
    routine_codes = [
        code
        for code, flag in (
            ('labels', routine.labels.exists()),
            ('fit_in_week', routine.fit_in_week),
        )
        if flag
    ]
    days = list(
        routine.days.order_by('order', 'id').prefetch_related(
            'slots__entries__exercise',
            'slots__entries__repetition_unit',
            'slots__entries__weight_unit',
            *[f'slots__entries__{name}' for name in CONFIG_SETS],
        )
    )
    names = _exercise_names(
        {e.exercise_id for d in days for s in d.slots.all() for e in s.entries.all()}, language
    )

    rows = []
    for day_order, day in enumerate(days, start=1):
        day_codes = _codes(
            day, ('day_type', day.type != 'custom'), ('need_logs', day.need_logs_to_advance)
        )
        day_row = {
            'routine_id': routine.pk,
            'routine_name': routine.name,
            'routine_description': routine.description,
            'start': routine.start.isoformat(),
            'end': routine.end.isoformat(),
            'day_id': day.pk,
            'day_order': day_order,
            'day_name': day.name,
            'day_description': day.description,
            'day_is_rest': 'yes' if day.is_rest else 'no',
        }
        slots = sorted(day.slots.all(), key=lambda s: (s.order, s.pk))
        if day.is_rest or not slots:
            rows.append({**day_row, 'unsupported': ';'.join(routine_codes + day_codes)})
            continue

        for slot_order, slot in enumerate(slots, start=1):
            slot_row = {
                **day_row,
                'slot_id': slot.pk,
                'slot_order': slot_order,
                'slot_comment': slot.comment,
            }
            slot_codes = routine_codes + day_codes + _codes(slot)
            entries = sorted(slot.entries.all(), key=lambda e: (e.order, e.pk))
            if not entries:
                rows.append({**slot_row, 'unsupported': ';'.join(slot_codes)})

            for entry_order, entry in enumerate(entries, start=1):
                configs = {name: list(getattr(entry, name).all()) for name in CONFIG_SETS}
                every = [c for group in configs.values() for c in group]
                entry_codes = _codes(
                    entry,
                    ('progression', any(c.iteration > 1 for c in every)),
                    ('requirements', any(c.requirements for c in every)),
                    ('custom_class', bool(entry.class_name)),
                    ('max_sets', bool(configs['maxsetsconfig_set'])),
                    (
                        'rounding',
                        entry.repetition_rounding != profile.repetitions_rounding
                        or entry.weight_rounding != profile.weight_rounding,
                    ),
                )
                row = {
                    **slot_row,
                    'entry_id': entry.pk,
                    'entry_order': entry_order,
                    'exercise_id': entry.exercise_id,
                    'exercise_uuid': str(entry.exercise.uuid),
                    'exercise_name': names.get(entry.exercise_id, ''),
                    'set_type': entry.type,
                    'rep_unit': entry.repetition_unit.name if entry.repetition_unit else '',
                    'weight_unit': entry.weight_unit.name if entry.weight_unit else '',
                    'notes': entry.comment,
                    'unsupported': ';'.join(slot_codes + entry_codes),
                }
                for column, model in CONFIG_COLUMNS.items():
                    base = [c for c in configs[f'{model.__name__.lower()}_set'] if c.iteration == 1]
                    if base:
                        value = Decimal(base[0].value)
                        row[column] = int(value) if BOUNDS.get(column, (0, 0, False))[2] else value
                rows.append(row)

    return [[row.get(column) for column in COLUMNS] for row in rows]


def _csv_cell(value) -> str:
    if value is None:
        return ''
    if isinstance(value, Decimal):
        return _fmt(value)
    return escape(value) if isinstance(value, str) else str(value)


def write_csv(rows: list[list]) -> bytes:
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(COLUMNS)
    writer.writerows([_csv_cell(value) for value in row] for row in rows)
    return out.getvalue().encode('utf-8-sig')


def _help_rows() -> list[list[str]]:
    """The XLSX help sheet: how to use the file, then one row per column"""
    rep_units = ', '.join(RepetitionUnit.objects.order_by('pk').values_list('name', flat=True))
    weight_units = ', '.join(WeightUnit.objects.order_by('pk').values_list('name', flat=True))
    rir = ', '.join(str(o) for o in RIR_OPTIONS[1:])
    days = Routine.MAX_DURATION_DAYS
    ids = 'routine_id, day_id, slot_id and entry_id'
    keep_id = 'Leave as exported. Blank on a new row. Clear it when copying into a new routine.'
    edit = 'Edit.'
    clears = 'Edit. Blank removes it on update.'

    guide = [
        'HOW TO USE THIS FILE',
        'Only the first sheet (routine) is imported. This help sheet is ignored. '
        'Keep the header row of the routine sheet exactly as it is.',
        '',
        'COLOURS',
        f'Red columns ({", ".join(PROTECTED_COLUMNS)}) mean do not edit: export fills them to '
        'tie each row to what already exists in the routine. Leave them as exported. '
        'Exceptions: a new row leaves them blank, and copying into a new routine clears them '
        'on every row (unsupported may stay if you choose to leave those settings out). See '
        'UPDATING AND COPYING.',
        'Blue columns are yours to edit. The COLUMNS list below says how for each one.',
        '',
        'GETTING STARTED',
        '1. Export the routine as XLSX, or download the blank template for a new routine.',
        '2. Edit the routine sheet. Example: to turn 3 sets of 8 into 4 sets of 6-8 reps, '
        'set sets 4, reps 6 and max_reps 8 on that row.',
        '3. Save as .xlsx (or .csv) and import it: update for the routine you exported, '
        'create for a brand-new routine.',
        '4. Check the preview: every row, which exercise each name matched, and what would be '
        'created, changed or deleted. Nothing is saved until you confirm.',
        '',
        'HOW ROWS WORK',
        'One row is one exercise in the plan, not one set you performed. '
        '4 sets of squats is one row with sets 4.',
        'Rows with the same day_order are one day. Days run in day_order order (1, 2, 3). '
        'Gaps are fine; the next export numbers them 1, 2, 3 again.',
        'Inside a day, rows with the same slot_order are one slot, and slots run in '
        'slot_order order. Two or more exercises in the same slot are a superset, done back '
        'to back in entry_order order.',
        'A rest day is one row with day_is_rest yes and no slot or exercise. '
        'An empty day or empty slot is one row with no exercise.',
        '',
        'UPDATING AND COPYING',
        f'Leave {ids} exactly as exported. They say which existing day, slot or exercise '
        'row to change. A new row has them blank.',
        'A row you delete is removed from the routine. Anything with logged workouts cannot '
        'be removed; the preview says so. Logged workouts are never changed.',
        'A blank cell does not mean "keep the old value". On update a blank reps, max_reps, '
        'weight, max_weight, rir, max_rir, rest, max_rest, notes or slot_comment removes it '
        '(a value with a progression cannot be cleared). Blank rep_unit means Repetitions, '
        'blank weight_unit kg, blank set_type normal.',
        f'To copy into a new routine, clear {ids} on every row, clear unsupported (or choose '
        'to leave those settings out) and give it a new routine_name. Keep exercise_id and '
        'exercise_uuid: they say which exercise it is, not where it sits in the routine.',
        'Changing an exercise: exercise_id is used first, then exercise_uuid, then '
        'exercise_name. Either clear exercise_id and exercise_uuid and type the new '
        "exercise_name, or put the new exercise's id and name and clear exercise_uuid. "
        'A new name next to the old id or uuid is refused as a mismatch.',
        '',
        'LIMITS',
        'There is no AMRAP or timer column. Write AMRAP in notes, e.g. "Last set AMRAP". '
        'rest and timed reps are planned numbers only; the file does not start timers.',
        'Values that start with = + - @ must be prefixed with an apostrophe. Formulas are refused.',
        '',
        'COLUMNS',
    ]
    columns = {
        'routine_id': (
            'Number of the routine this file came from, filled by export, e.g. 24.',
            'Leave as exported to update. Clear it to create a new routine.',
        ),
        'routine_name': (
            'Routine name, at most 25 characters, e.g. Upper Lower. Required on the first '
            'row; later rows blank or the same.',
            edit,
        ),
        'routine_description': (
            'Text about the routine, at most 1000 characters. First row; later rows blank or '
            'the same.',
            edit,
        ),
        'start': ('First date of the routine, YYYY-MM-DD, e.g. 2026-10-05. First row.', edit),
        'end': (
            f'Last date, YYYY-MM-DD, not before start and at most {days} days after it. First row.',
            edit,
        ),
        'day_id': ('Number of an existing day, filled by export.', keep_id),
        'day_order': (
            'Which day the row belongs to and where that day goes: whole number, 1 = first. '
            'Every row needs one.',
            'Edit to add, reorder or move days.',
        ),
        'day_name': (
            'Day name, at most 20 characters, e.g. Upper A. Same on every row of the day.',
            edit,
        ),
        'day_description': (
            'Text about the day, at most 1000 characters. Same on every row of the day.',
            edit,
        ),
        'day_is_rest': ('yes for a rest day; no or blank for a training day.', edit),
        'slot_id': ('Number of an existing slot, filled by export.', keep_id),
        'slot_order': (
            'Position of the slot in the day: whole number, 1 = first. Required on exercise '
            'rows. Same day_order and slot_order = same slot = superset.',
            'Edit to reorder or make supersets.',
        ),
        'slot_comment': (
            'Note for the whole slot, at most 200 characters, e.g. No rest inside the '
            'superset. Same on every row of the slot.',
            clears,
        ),
        'entry_id': ('Number of an existing exercise row of the plan, filled by export.', keep_id),
        'entry_order': (
            'Order of the exercises inside a slot: whole number, 1 = first, used once per '
            'slot. Required.',
            edit,
        ),
        'exercise_id': (
            "The exercise's wger number, e.g. 73. Says which exercise it is and is checked "
            'first. Not a place in the routine, so it stays when copying.',
            'Leave as exported. To change the exercise, see UPDATING AND COPYING.',
        ),
        'exercise_uuid': (
            "The exercise's permanent code, used when exercise_id is blank. Stays when copying.",
            'Leave as exported. To change the exercise, see UPDATING AND COPYING.',
        ),
        'exercise_name': (
            'Exact exercise name or alias (any capitals), in your language or English, e.g. '
            'Bench Press. Next to an id or uuid it must belong to that exercise.',
            'Type it on new rows. The preview shows the match.',
        ),
        'set_type': (f'Kind of set: {", ".join(ExerciseType.values)}. Blank = normal.', edit),
        'sets': ('How many sets of this exercise: whole number 1-50, e.g. 4. Required.', edit),
        'reps': (
            'Target per set, counted in rep_unit: 0-3000, up to 2 decimals. e.g. 8 '
            'repetitions, or 45 with rep_unit Seconds for a 45 second hold.',
            clears,
        ),
        'max_reps': (
            'Top of a reps range: reps 8 and max_reps 12 means 8-12. A plan target, not a '
            'personal record. Not below reps.',
            clears,
        ),
        'rep_unit': (
            f'Unit of reps and max_reps: {rep_units}. Blank = Repetitions. For planned time '
            'use Seconds and type plain seconds in reps, e.g. 45, not 0:45 or 45s.',
            edit,
        ),
        'weight': (
            'Target per set, counted in weight_unit: 0-3000, up to 2 decimals, e.g. 60 (kg). '
            'With Kilometers Per Hour it is a speed, e.g. 10 = 10 km/h.',
            clears,
        ),
        'max_weight': (
            'Top of a weight range: weight 60 and max_weight 70 means 60-70. A plan target, '
            'not a personal record. Not below weight.',
            clears,
        ),
        'weight_unit': (
            f'Unit of weight and max_weight: {weight_units}. Blank = kg. kg is a load; use '
            'Kilometers Per Hour, not kg, when weight is a speed.',
            edit,
        ),
        'rir': (
            f'Reps in reserve, the reps left in the tank after each set: one of {rir}, e.g. 2.',
            clears,
        ),
        'max_rir': (
            'Top of a reps in reserve range, e.g. rir 1 and max_rir 2. Not below rir.',
            clears,
        ),
        'rest': (
            'Rest between sets in seconds: whole number 0-1800, e.g. 120 = 2 minutes.',
            clears,
        ),
        'max_rest': (
            'Top of a rest range in seconds: whole number 0-600, e.g. rest 90 and max_rest 120. '
            'Not below rest.',
            clears,
        ),
        'notes': (
            'Note for this exercise, at most 100 characters, e.g. Pause at the bottom, or '
            'Last set AMRAP.',
            clears,
        ),
        'unsupported': (
            'Filled by export: settings this file cannot show, separated by ;, e.g. '
            'progression. Update keeps them as they are. Create stops unless you choose to '
            'leave them out.',
            'Leave as exported.',
        ),
    }
    return [
        *([line] for line in guide),
        ['column', 'what it means', 'do you edit it?'],
        *([column, *columns[column]] for column in COLUMNS),
    ]


PROTECTED_COLUMNS = ('routine_id', 'day_id', 'slot_id', 'entry_id', 'unsupported')
"""Filled by export and matched by the importer, shown red in XLSX"""

_RED = openpyxl.styles.PatternFill('solid', fgColor='C00000')
_RED_VALUE = openpyxl.styles.PatternFill('solid', fgColor='F8CBAD')
_BLUE = openpyxl.styles.PatternFill('solid', fgColor='DDEBF7')
_WHITE_BOLD = openpyxl.styles.Font(bold=True, color='FFFFFF')
_BOLD = openpyxl.styles.Font(bold=True)


def write_xlsx(rows: list[list], with_help: bool = False) -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'routine'
    sheet.append(COLUMNS)
    for cell in sheet[1]:
        protected = cell.value in PROTECTED_COLUMNS
        cell.fill = _RED if protected else _BLUE
        cell.font = _WHITE_BOLD if protected else _BOLD
    for number, row in enumerate(rows, start=2):
        for column, value in enumerate(row, start=1):
            if value is None:
                continue
            cell = sheet.cell(number, column, escape(value) if isinstance(value, str) else value)
            if isinstance(value, str):
                # Always a string, openpyxl must never store a formula
                cell.data_type = 's'
            if COLUMNS[column - 1] in PROTECTED_COLUMNS:
                cell.fill = _RED_VALUE

    if with_help:
        help_sheet = workbook.create_sheet('help')
        for letter, width in zip('ABC', (20, 100, 60)):
            help_sheet.column_dimensions[letter].width = width
        for line in _help_rows():
            help_sheet.append(line)
            for cell in help_sheet[help_sheet.max_row]:
                cell.data_type = 's'
            if line[0] in PROTECTED_COLUMNS and len(line) == 3:
                for cell in help_sheet[help_sheet.max_row]:
                    cell.fill = _RED
                    cell.font = _WHITE_BOLD

    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def download(rows: list[list], kind: str, name: str, with_help: bool = False) -> HttpResponse:
    if kind == 'xlsx':
        response = HttpResponse(write_xlsx(rows, with_help), content_type=XLSX_TYPE)
    else:
        response = HttpResponse(write_csv(rows), content_type=CSV_TYPE)
    response['Content-Disposition'] = f'attachment; filename="{name}.{kind}"'
    return response
